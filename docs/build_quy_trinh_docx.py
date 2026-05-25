from pathlib import Path
import re

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_CELL_VERTICAL_ALIGNMENT
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt, RGBColor


ROOT = Path(__file__).resolve().parents[1]
PROCESS_MD = ROOT / "docs" / "QUY_TRINH_NHAN_BAN_LOVABLE_SUPABASE_CLOUDFLARE.md"
SKILL_MD = ROOT / "skills" / "lovable-supabase-cloudflare-admin" / "SKILL.md"
OUT = ROOT / "docs" / "Quy_Trinh_Nhan_Ban_Lovable_Supabase_Cloudflare_Admin.docx"


def set_cell_shading(cell, fill):
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:fill"), fill)
    tc_pr.append(shd)


def set_cell_border(cell, color="DADCE0", size="8"):
    tc = cell._tc
    tc_pr = tc.get_or_add_tcPr()
    borders = tc_pr.first_child_found_in("w:tcBorders")
    if borders is None:
        borders = OxmlElement("w:tcBorders")
        tc_pr.append(borders)
    for edge in ("top", "left", "bottom", "right"):
        tag = "w:{}".format(edge)
        element = borders.find(qn(tag))
        if element is None:
            element = OxmlElement(tag)
            borders.append(element)
        element.set(qn("w:val"), "single")
        element.set(qn("w:sz"), size)
        element.set(qn("w:space"), "0")
        element.set(qn("w:color"), color)


def style_document(doc):
    section = doc.sections[0]
    section.top_margin = Cm(2.2)
    section.bottom_margin = Cm(2.2)
    section.left_margin = Cm(2.2)
    section.right_margin = Cm(2.2)

    styles = doc.styles
    normal = styles["Normal"]
    normal.font.name = "Arial"
    normal._element.rPr.rFonts.set(qn("w:eastAsia"), "Arial")
    normal.font.size = Pt(10.5)
    normal.paragraph_format.space_after = Pt(6)
    normal.paragraph_format.line_spacing = 1.15

    for name, size, color, before, after in [
        ("Heading 1", 18, "0F172A", 14, 6),
        ("Heading 2", 14, "1F4D78", 12, 5),
        ("Heading 3", 12, "334155", 10, 4),
    ]:
        style = styles[name]
        style.font.name = "Arial"
        style._element.rPr.rFonts.set(qn("w:eastAsia"), "Arial")
        style.font.size = Pt(size)
        style.font.bold = True
        style.font.color.rgb = RGBColor.from_string(color)
        style.paragraph_format.space_before = Pt(before)
        style.paragraph_format.space_after = Pt(after)
        style.paragraph_format.keep_with_next = True


def add_title(doc):
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_after = Pt(4)
    run = p.add_run("Quy Trình Nhân Bản Website")
    run.font.name = "Arial"
    run._element.rPr.rFonts.set(qn("w:eastAsia"), "Arial")
    run.font.size = Pt(24)
    run.font.bold = True
    run.font.color.rgb = RGBColor.from_string("0F172A")

    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_after = Pt(14)
    run = p.add_run("Lovable -> GitHub -> Cloudflare -> Supabase -> Dashboard Admin")
    run.font.name = "Arial"
    run._element.rPr.rFonts.set(qn("w:eastAsia"), "Arial")
    run.font.size = Pt(11)
    run.font.color.rgb = RGBColor.from_string("475569")

    add_callout(
        doc,
        "Ghi nhớ quan trọng",
        "Mọi tính năng admin như thêm học viên, thêm/sửa/xóa bài giảng phải chạy qua server và cần đủ 3 biến Supabase: SUPABASE_URL, SUPABASE_PUBLISHABLE_KEY, SUPABASE_SERVICE_ROLE_KEY.",
        fill="FFF7E6",
        border="F59E0B",
    )


def add_callout(doc, title, body, fill="F8FAFC", border="CBD5E1"):
    table = doc.add_table(rows=1, cols=1)
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.autofit = True
    cell = table.cell(0, 0)
    cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
    set_cell_shading(cell, fill)
    set_cell_border(cell, color=border, size="10")
    p = cell.paragraphs[0]
    p.paragraph_format.space_after = Pt(3)
    r = p.add_run(title)
    r.bold = True
    r.font.name = "Arial"
    r._element.rPr.rFonts.set(qn("w:eastAsia"), "Arial")
    r.font.color.rgb = RGBColor.from_string("0F172A")
    p2 = cell.add_paragraph(body)
    p2.paragraph_format.space_after = Pt(0)
    doc.add_paragraph()


def add_code_block(doc, lines):
    text = "\n".join(lines).strip()
    if not text:
        return
    table = doc.add_table(rows=1, cols=1)
    cell = table.cell(0, 0)
    set_cell_shading(cell, "F1F5F9")
    set_cell_border(cell, color="CBD5E1", size="6")
    p = cell.paragraphs[0]
    p.paragraph_format.space_before = Pt(3)
    p.paragraph_format.space_after = Pt(3)
    for idx, line in enumerate(text.splitlines()):
        if idx:
            p.add_run().add_break()
        run = p.add_run(line)
        run.font.name = "Consolas"
        run._element.rPr.rFonts.set(qn("w:eastAsia"), "Consolas")
        run.font.size = Pt(9)
        run.font.color.rgb = RGBColor.from_string("0F172A")
    doc.add_paragraph()


def add_markdown_inline(paragraph, text):
    parts = re.split(r"(`[^`]+`|\*\*[^*]+\*\*)", text)
    for part in parts:
        if not part:
            continue
        if part.startswith("`") and part.endswith("`"):
            run = paragraph.add_run(part[1:-1])
            run.font.name = "Consolas"
            run._element.rPr.rFonts.set(qn("w:eastAsia"), "Consolas")
            run.font.size = Pt(9.5)
            run.font.color.rgb = RGBColor.from_string("B45309")
        elif part.startswith("**") and part.endswith("**"):
            run = paragraph.add_run(part[2:-2])
            run.bold = True
        else:
            paragraph.add_run(part)


def add_markdown(doc, md_text, heading_prefix=None):
    in_code = False
    code_lines = []
    in_frontmatter = False

    for raw in md_text.splitlines():
        line = raw.rstrip()

        if line.strip() == "---" and not in_code:
            in_frontmatter = not in_frontmatter
            continue
        if in_frontmatter:
            continue

        if line.strip().startswith("```"):
            if not in_code:
                in_code = True
                code_lines = []
            else:
                add_code_block(doc, code_lines)
                in_code = False
            continue

        if in_code:
            code_lines.append(line)
            continue

        if not line.strip():
            continue

        if line.startswith("# "):
            text = line[2:].strip()
            if heading_prefix:
                text = f"{heading_prefix}: {text}"
            doc.add_heading(text, level=1)
            continue
        if line.startswith("## "):
            doc.add_heading(line[3:].strip(), level=2)
            continue
        if line.startswith("### "):
            doc.add_heading(line[4:].strip(), level=3)
            continue

        numbered = re.match(r"^(\d+)\.\s+(.*)$", line)
        bullet = re.match(r"^-\s+(.*)$", line)
        checkbox = re.match(r"^\[ \]\s+(.*)$", line)

        if numbered:
            p = doc.add_paragraph(style="List Number")
            add_markdown_inline(p, numbered.group(2))
        elif bullet:
            p = doc.add_paragraph(style="List Bullet")
            add_markdown_inline(p, bullet.group(1))
        elif checkbox:
            p = doc.add_paragraph(style="List Bullet")
            add_markdown_inline(p, "☐ " + checkbox.group(1))
        else:
            p = doc.add_paragraph()
            add_markdown_inline(p, line)


def add_footer(doc):
    section = doc.sections[0]
    footer = section.footer
    p = footer.paragraphs[0]
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = p.add_run("Quy trình nhân bản Lovable + Supabase + Cloudflare")
    run.font.name = "Arial"
    run._element.rPr.rFonts.set(qn("w:eastAsia"), "Arial")
    run.font.size = Pt(8)
    run.font.color.rgb = RGBColor.from_string("64748B")


def main():
    doc = Document()
    style_document(doc)
    add_title(doc)

    doc.add_heading("Phần 1: Quy Trình Chuẩn", level=1)
    add_markdown(doc, PROCESS_MD.read_text(encoding="utf-8"))

    doc.add_page_break()
    doc.add_heading("Phần 2: Skill Dùng Lại Cho Codex", level=1)
    add_callout(
        doc,
        "Cách dùng skill",
        "Khi nhân bản hoặc debug dự án tương tự, đưa nội dung skill này cho Codex để nó đi đúng quy trình kiểm tra Lovable, GitHub, Cloudflare, Supabase và dashboard admin.",
        fill="EEF6FF",
        border="60A5FA",
    )
    add_markdown(doc, SKILL_MD.read_text(encoding="utf-8"), heading_prefix="Skill")

    add_footer(doc)
    doc.save(OUT)
    print(OUT)


if __name__ == "__main__":
    main()

